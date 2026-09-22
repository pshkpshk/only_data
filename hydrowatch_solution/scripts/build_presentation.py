"""Собрать docs/presentation.pptx (10 слайдов) из плана PRESENTATION.md.

Запуск:  python scripts/build_presentation.py [--screenshots docs/screenshots]
Если в каталоге скриншотов есть 01_map_swipe.png / 02_report_panel.png / 03_swagger.png,
они вставляются на слайды 7 и 9; иначе остаются рамки-подсказки.
Зависимость: python-pptx (pip install python-pptx).
"""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x0B, 0x1F, 0x3A)
BLUE = RGBColor(0x1F, 0x77, 0xB4)
RED = RGBColor(0xE6, 0x39, 0x46)
GREY = RGBColor(0x55, 0x5F, 0x6B)
LIGHT = RGBColor(0xF3, 0xF6, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

W, H = Inches(13.333), Inches(7.5)


def add_bg(slide, color):
    shp = slide.shapes.add_shape(1, 0, 0, W, H)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def add_text(slide, left, top, width, height, text, size=18, bold=False, color=NAVY,
             align=PP_ALIGN.LEFT, font="Calibri"):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
    return tb


def add_bullets(slide, left, top, width, height, items, size=17, color=NAVY):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        level = 0
        if item.startswith("  "):
            level = 1
            item = item.strip()
        p.level = level
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = ("• " if level == 0 else "– ") + item
        r.font.size = Pt(size - 2 * level)
        r.font.color.rgb = color if level == 0 else GREY
        r.font.name = "Calibri"
    return tb


def header(slide, number, title, subtitle=None):
    add_bg(slide, WHITE)
    bar = slide.shapes.add_shape(1, 0, 0, W, Inches(1.05))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    add_text(slide, Inches(0.5), Inches(0.22), Inches(11.5), Inches(0.7), f"{number}. {title}",
             size=28, bold=True, color=WHITE)
    add_text(slide, Inches(11.6), Inches(0.3), Inches(1.5), Inches(0.5), "HydroWatch Amur",
             size=12, color=RGBColor(0xBF, 0xD7, 0xEA), align=PP_ALIGN.RIGHT)
    if subtitle:
        add_text(slide, Inches(0.5), Inches(1.15), Inches(12.3), Inches(0.5), subtitle,
                 size=15, color=GREY)


def add_table(slide, left, top, width, rows, col_widths=None, size=13, header_fill=BLUE):
    n_rows, n_cols = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, Inches(0.4) * n_rows)
    table = shape.table
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = w
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = str(value)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(size)
                    run.font.name = "Calibri"
                    run.font.bold = r == 0
                    run.font.color.rgb = WHITE if r == 0 else NAVY
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r == 0 else (LIGHT if r % 2 else WHITE)
    return shape


def add_picture_or_placeholder(slide, path: Path | None, left, top, width, height, hint):
    if path and path.is_file():
        pic = slide.shapes.add_picture(str(path), left, top)
        # вписать в рамку с сохранением пропорций
        ratio = min(width / pic.width, height / pic.height)
        pic.width, pic.height = Emu(int(pic.width * ratio)), Emu(int(pic.height * ratio))
        pic.left, pic.top = left, top
        return pic
    box = slide.shapes.add_shape(1, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT
    box.line.color.rgb = GREY
    box.line.dash_style = 4
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = hint
    r.font.size = Pt(14)
    r.font.color.rgb = GREY
    return box


def build(out: Path, shots: Path, team: str) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    # 1 -------------------------------------------------------------- title / задача
    s = prs.slides.add_slide(blank)
    add_bg(s, NAVY)
    add_text(s, Inches(0.8), Inches(1.4), Inches(11.5), Inches(1.2), "HydroWatch Amur",
             size=48, bold=True, color=WHITE)
    add_text(s, Inches(0.8), Inches(2.5), Inches(11.5), Inches(1.0),
             "Динамика затопления по Sentinel-1 + Sentinel-2: вода «до» и «пик», зона нового "
             "затопления, площади и сервис", size=22, color=RGBColor(0xBF, 0xD7, 0xEA))
    add_bullets(s, Inches(0.8), Inches(3.7), Inches(11.5), Inches(2.6), [
        "Вход: пара наблюдений «до / пик» (SAR VV,VH + опциональная оптика + рельеф/GSW). "
        "Выход: маски воды на две даты, маска нового затопления, площади в га",
        "11 пар, 5 районов Амурской области, паводки 2019/2021 и межень 2018",
        "Метрика: 0.45·Q_flood + 0.25·Q_peak + 0.15·Q_pre + 0.15·Spec_межень",
        "Всепогодность SAR + точность оптики; сервис с картой, отчётом и REST API",
    ], size=18, color=WHITE)
    add_text(s, Inches(0.8), Inches(6.6), Inches(11.5), Inches(0.5),
             f"КосмоХакатон 2026 · кейс «Гидрологический мониторинг» · команда {team}",
             size=14, color=RGBColor(0xBF, 0xD7, 0xEA))

    # 2 -------------------------------------------------------------- данные
    s = prs.slides.add_slide(blank)
    header(s, 2, "Что в данных", "EDA по 11 парам (REPORT §2)")
    add_bullets(s, Inches(0.5), Inches(1.8), Inches(6.6), Inches(5), [
        "Зона нового затопления — всего 0.02–2 % площади района: сильный дисбаланс классов",
        "Пригодной оптики нет у 5 из 11 пар (облачность) → решение обязано работать на одном SAR",
        "Разрыв SAR–оптика до 5 суток; между «до» и «пик» — ~6 недель: разность масок ловит сезонность",
        "HAND: медиана 0.5 м под водой против 3.7 м на фоне — рельеф информативен, но не как жёсткий фильтр",
        "σ⁰: VV бимодальна (вода ≈ −20…−24 дБ), VH ниже на 6–8 дБ и с меньшим контрастом",
    ], size=17)
    add_table(s, Inches(7.4), Inches(1.9), Inches(5.5), [
        ["Пара", "flood, га", "% AOI", "оптика"],
        ["Зея-2021 Свободный", "2484", "1.9", "да"],
        ["Поярково-2021", "2167", "1.6", "да"],
        ["Благовещенск-2021", "883", "0.7", "да"],
        ["Константиновка-2019", "644", "0.5", "нет"],
        ["Белогорск-2019", "187", "0.1", "да"],
        ["Межень-2018 (3 пары)", "25–197", "≤0.2", "нет"],
    ], col_widths=[Inches(2.5), Inches(1.0), Inches(0.9), Inches(1.1)], size=12)

    # 3 -------------------------------------------------------------- почему не Оцу
    s = prs.slides.add_slide(blank)
    header(s, 3, "Почему не «Оцу по VV + разность масок»", "Ограничения классических подходов (REPORT §4)")
    add_table(s, Inches(0.5), Inches(1.9), Inches(12.3), [
        ["Источник", "Где ошибается", "Следствие для площади"],
        ["SAR, порог по VV", "ветровая рябь, радиотени, гладкие сухие поверхности (пашня, асфальт)",
         "ложная «вода» на суше, пропуски на волнующейся воде"],
        ["Оптика, NDWI/MNDWI", "мутная паводковая вода, тени облаков, отсутствие снимка",
         "недооценка пика, дыры в покрытии"],
        ["Разность масок за 6 недель", "сезонный ход уровня, разные орбиты и углы",
         "«затопление» там, где просто изменился уровень"],
    ], col_widths=[Inches(2.6), Inches(5.2), Inches(4.5)], size=14)
    add_bullets(s, Inches(0.5), Inches(4.6), Inches(12.3), Inches(2.5), [
        "Ошибки SAR и оптики независимы → объединять надо на уровне признаков, а не голосованием масок",
        "Затопление — отдельная цель: сеть учится сразу трём картам (вода «до», вода «пик», прирост) "
        "с гидрологической связностью, а не арифметике порогов",
        "Baseline «Оцу + разность» оставлен как нижняя граница шкалы в отчёте",
    ], size=17)

    # 4 -------------------------------------------------------------- схема решения
    s = prs.slides.add_slide(blank)
    header(s, 4, "Схема решения", "Siamese multimodal U-Net: два энкодера, temporal-признаки, три головы")
    diagram = [
        "S1 pre/peak (VV,VH,|Δ|) ──► ResNet-18 (SAR) ─┐",
        "S2 pre/peak (B03,B04,B08,B11) ► ResNet-18 (optical, SSL4EO-S12 MoCo) ─┤  → temporal-признаки pre / peak / Δ",
        "AUX: HAND, уклон, GSW occurrence, WorldCover, флаг оптики ──────────────┘        на каждом масштабе",
        "                                   │",
        "                          U-Net decoder (skip-connections)",
        "                                   │",
        "        ┌──────────────────────────┼──────────────────────────┐",
        "   вода «до» (σ)            вода «пик» (σ)          новое затопление (σ)",
        "        └── пороги 0.95 / 0.78 / 0.95 (калибровка на отложенной Зее-2021) ──┘",
    ]
    box = s.shapes.add_shape(1, Inches(0.5), Inches(1.8), Inches(12.3), Inches(3.3))
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT
    box.line.color.rgb = GREY
    tf = box.text_frame
    tf.word_wrap = False
    for i, line in enumerate(diagram):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = line
        r.font.size = Pt(11)
        r.font.name = "Menlo"
        r.font.color.rgb = NAVY
    add_bullets(s, Inches(0.5), Inches(5.25), Inches(12.3), Inches(2.0), [
        "Modality dropout 50 %: оптика случайно выключается при обучении → одна модель для пар с оптикой и без",
        "Loss = focal Tversky + Dice/BCE + штраф за площадь + гидро-логика (flood ⊂ peak, flood ∩ pre = ∅)",
        "Инференс окнами 384×384 (stride 256), AMP; веса best.pt (SHA-256 6c7ef209…2d06)",
    ], size=16)

    # 5 -------------------------------------------------------------- эксперименты
    s = prs.slides.add_slide(blank)
    header(s, 5, "Эксперименты и результат", "Валидация по событию: Зея-2021 отложена целиком, seed 2026")
    add_table(s, Inches(0.5), Inches(1.85), Inches(6.4), [
        ["Эксперимент", "val loss"],
        ["Базовая модель с нуля (исправленные сцены)", "0.614"],
        ["Fine-tune лучшей ранней модели, lr 1e-5, 6 эп.", "0.489 → 0.449"],
        ["Калибровка порогов 0.95 / 0.78 / 0.95", "Dice + площадь ↑"],
    ], col_widths=[Inches(4.6), Inches(1.8)], size=13)
    add_table(s, Inches(7.2), Inches(1.85), Inches(5.6), [
        ["Компонент метрики", "11 пар (self-check)", "Отложенная Зея"],
        ["Q_flood (0.45)", "0.699", "0.53"],
        ["Q_water_peak (0.25)", "0.768", "0.57"],
        ["Q_water_pre (0.15)", "0.680", "1.00"],
        ["Spec_base (0.15)", "1.000", "—"],
        ["Score", "0.758", "—"],
    ], col_widths=[Inches(2.2), Inches(1.9), Inches(1.5)], size=13)
    add_bullets(s, Inches(0.5), Inches(4.7), Inches(12.3), Inches(2.5), [
        "Честная цифра — отложенная Зея: q_flood 0.53 при 3211 га прогноза против 2179 га в JSON эталона "
        "(и 2484 га по самой маске эталона)",
        "Межень: flood = 0 по правилу события (event_kind), Spec_base = 1.0 без подгонки под эталон",
        "Абляции тем же чекпоинтом: --ablation no_optical / no_aux, пороги 0.5 — команды в REPORT §6.2",
    ], size=16)

    # 6 -------------------------------------------------------------- критика эталона
    s = prs.slides.add_slide(blank)
    header(s, 6, "Критика эталона и метрики", "scripts/reference_audit.py — воспроизводимый аудит выданных масок")
    add_bullets(s, Inches(0.5), Inches(1.8), Inches(12.3), Inches(5.3), [
        "5 из 11 пар: маска воды почти не содержит постоянной воды GSW — сцены орбиты 105 покрывали AOI частично",
        "  Константиновка-2021: «до» = 30 га при 5436 га постоянной воды реки; Благовещенск-2021: 2649 га при 7247 га",
        "Зея-2021: JSON эталона 2179 га, а по самой маске эталона — 2484 га (+14 %)",
        "flood эталона ≠ peak ∧ ¬pre ∧ ¬perm: расхождение до 1156 га (Благовещенск-2019) — фильтры уклон/HAND/MMU "
        "применены только к flood",
        "Наш случай: Поярково — сцены орбиты 105 покрывали 0.3–0.4 % AOI → configs/scene_overrides.yaml "
        "заменяет их на орбиту 32 (100 %)",
        "Вывод: часть «ошибки» модели — ошибка эталона; метрика по площадям не различает «где», только «сколько»",
    ], size=17)

    # 7 -------------------------------------------------------------- сервис
    s = prs.slides.add_slide(blank)
    header(s, 7, "Сервис: карта, отчёт, выгрузки", "docker compose up → http://localhost:8000 · REST API /docs")
    add_picture_or_placeholder(s, shots / "01_map_swipe.png", Inches(0.5), Inches(1.8),
                               Inches(7.0), Inches(4.3), "скриншот: карта со шторкой «до | пик»")
    add_picture_or_placeholder(s, shots / "02_report_panel.png", Inches(7.7), Inches(1.8),
                               Inches(5.1), Inches(4.3), "скриншот: панель отчёта")
    add_bullets(s, Inches(0.5), Inches(6.2), Inches(12.3), Inches(1.2), [
        "Запрос: полигон/bbox + даты → подбор пары наблюдений; слои «до / пик / прирост / убыль», шторка, bbox мышью; "
        "Зея-2021: +3211 га прироста, 640 га убыли, +71 % зеркала",
        "Отчёт JSON/CSV: га и км², доля от территории, разбивка по типам поверхности, профиль HAND; "
        "GeoJSON (EPSG:4326, id/type/area_ha), GeoTIFF",
    ], size=14)

    # 8 -------------------------------------------------------------- продукт
    s = prs.slides.add_slide(blank)
    header(s, 8, "Продуктовый сценарий и обновление по новым сценам", "Кто пользуется и что пересчитывается")
    add_bullets(s, Inches(0.5), Inches(1.8), Inches(6.3), Inches(5.3), [
        "Пользователь: дежурная смена ЦУКС / ЦГМС — решение нужно в часы после пролёта",
        "Хранится: опорное состояние «до» (маска воды + дата + орбита), постоянная вода GSW, AUX-слои района",
        "Новая S1-сцена → POST /api/observations: маска воды сравнивается с опорным состоянием, "
        "пересчитывается только разность и отчёт",
        "Разные орбиты (32 / 105) — отдельные ряды наблюдений, чтобы не путать геометрию съёмки с изменением",
        "Задержка: минуты на район 4000×4000; пересчёт «истории» не требуется",
    ], size=17)
    add_table(s, Inches(7.1), Inches(1.9), Inches(5.7), [
        ["Шаг", "Что делает", "Время"],
        ["1", "Пролёт S1 → выгрузка сцены на сетку AOI", "≈ 10 мин"],
        ["2", "hydrowatch-predict (маски)", "10–20 с (GPU)"],
        ["3", "POST /api/observations", "< 1 с"],
        ["4", "Отчёт, GeoJSON, карта", "≈ 1 с"],
    ], col_widths=[Inches(0.6), Inches(3.5), Inches(1.6)], size=13)

    # 9 -------------------------------------------------------------- воспроизводимость
    s = prs.slides.add_slide(blank)
    header(s, 9, "Воспроизводимость и ресурсы", "Одна команда на инференс, одна — на сервис")
    add_bullets(s, Inches(0.5), Inches(1.8), Inches(6.6), Inches(5.3), [
        "docker compose up --build → сервис; docker compose --profile inference run inference → submission.csv",
        "uv sync --all-extras / pip install -e .; зависимости зафиксированы: pyproject.toml + uv.lock",
        "seed 2026 везде; SHA-256 чекпоинта и всех артефактов (SHA256SUMS.txt); 22 теста, ruff",
        "Инференс: окна 384/256, batch 4, AMP — ≈ 2–3 ГБ VRAM, 10–20 с на пару; CPU — минуты на пару",
        "Сервис: < 1 ГБ RAM, отчёт по паре ≈ 1 с на 2 vCPU; без датасета стартует (без разбивки по покрову)",
        "Веса и маски — по ссылкам из README, доступ ≥ 10 рабочих дней",
    ], size=16)
    add_picture_or_placeholder(s, shots / "03_swagger.png", Inches(7.4), Inches(1.8),
                               Inches(5.4), Inches(4.9), "скриншот: Swagger /docs или docker compose up")

    # 10 ------------------------------------------------------------- ограничения
    s = prs.slides.add_slide(blank)
    header(s, 10, "Ограничения и развитие", "Что не решено и что делать дальше")
    add_bullets(s, Inches(0.5), Inches(1.8), Inches(6.2), Inches(5.3), [
        "Локальный угол падения не используется — сцены разных орбит нормализуются только статистически",
        "Затопленная растительность (double-bounce) не картируется как вода",
        "Переоценка на широких поймах (Поярково: 3367 против 2167 га) — размытая граница вода/влажная почва",
        "Пороги калиброваны на одном отложенном событии",
    ], size=17, color=RED)
    add_bullets(s, Inches(6.9), Inches(1.8), Inches(6.0), Inches(5.3), [
        "Временной ряд S1 вместо пары: устойчивая опорная вода и сезонная норма",
        "Расширение выборки событиями 2013 и 2022 гг. (Амур, Зея, Бурея)",
        "Слой double-bounce (VV/VH + текстура) для затопленной растительности",
        "Агрегация по бассейнам и водомерным постам, связка с уровнями Росгидромета",
        "Инкрементальный контур обновления — в продукт",
    ], size=17, color=BLUE)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"saved {out} ({out.stat().st_size / 1e6:.1f} MB), slides: {len(prs.slides)}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=Path("docs/presentation.pptx"))
    parser.add_argument("--screenshots", type=Path, default=Path("docs/screenshots"))
    parser.add_argument("--team", default="<название команды>")
    args = parser.parse_args()
    build(args.out, args.screenshots, args.team)


if __name__ == "__main__":
    main()
