"""Собрать docs/DEFENSE_MEMO.md в Памятка_защита.docx (простой markdown → docx)."""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt, RGBColor

SRC = Path(sys.argv[1] if len(sys.argv) > 1 else "docs/DEFENSE_MEMO.md")
OUT = Path(sys.argv[2] if len(sys.argv) > 2 else "Памятка_защита.docx")
BLUE, DARK = RGBColor(0x0A, 0x7A, 0xFF), RGBColor(0x0B, 0x3D, 0x91)

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2)
    s.top_margin = s.bottom_margin = Cm(1.8)
doc.styles["Normal"].font.name = "Helvetica Neue"
doc.styles["Normal"].font.size = Pt(11.5)


def add_runs(par, text):
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part:
            par.add_run(part)


table_rows, para_buf = [], []


def flush_table():
    global table_rows
    if not table_rows:
        return
    rows = [r for r in table_rows if not re.match(r"^\|[\s\-|:]+\|$", r)]
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.style = "Light Grid Accent 1"
    for ri, row in enumerate(cells):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.text = ""
            add_runs(cell.paragraphs[0], val)
            for r in cell.paragraphs[0].runs:
                r.font.size = Pt(10.5)
                r.bold = r.bold or ri == 0
    doc.add_paragraph()
    table_rows = []


def flush_para():
    global para_buf
    if para_buf:
        p = doc.add_paragraph()
        add_runs(p, " ".join(para_buf))
        p.paragraph_format.space_after = Pt(6)
        para_buf = []


for line in SRC.read_text().splitlines():
    if line.startswith("|"):
        flush_para()
        table_rows.append(line)
        continue
    flush_table()
    if line.startswith("# "):
        flush_para()
        for r in doc.add_heading(line[2:], level=0).runs:
            r.font.color.rgb = BLUE
    elif line.startswith("## "):
        flush_para()
        for r in doc.add_heading(line[3:], level=1).runs:
            r.font.color.rgb = BLUE
    elif line.startswith("### "):
        flush_para()
        for r in doc.add_heading(line[4:], level=2).runs:
            r.font.color.rgb = DARK
    elif line.strip() in ("---", ""):
        flush_para()
    elif re.match(r"^\s*- ", line):
        flush_para()
        add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*- ", "", line))
    elif re.match(r"^\d+\. ", line):
        flush_para()
        add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", line))
    elif para_buf or not doc.paragraphs or not line.startswith("  "):
        para_buf.append(line.strip())
    else:
        add_runs(doc.paragraphs[-1], " " + line.strip())
flush_para()
flush_table()
doc.save(OUT)
print(f"saved {OUT} ({OUT.stat().st_size // 1024} KB)")
